"""Tests of presentations read slide by slide, on decks made here with
python-pptx: python -m unittest test_slides (scripts/test.sh runs them).
They need docling's pptx reader (docling-slim[convert-core,format-pptx])
and the server's own packages, not the models of the image."""

import asyncio
import base64
import copy
import io
import os
import re
import sys
import unittest

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import server  # noqa: E402

FOOTER = "Allotment Society · Spring 2026"
LONG = "This first line runs on and on about soil, seeds and sheds, " * 4


def png(color, size=(40, 20)) -> io.BytesIO:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, format="PNG")
    buf.seek(0)
    return buf


def text_box(slide, text, top, size=18):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(8), Inches(0.6))
    box.text_frame.text = text
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
    return box


def slide_number(slide, n):
    """The layout's slide number placeholder on the slide, as PowerPoint puts
    it there (python-pptx leaves it out of a new slide)."""
    layout = [ph for ph in slide.slide_layout.placeholders if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER]
    assert layout, "the layout has no slide number"
    slide.shapes._spTree.append(copy.deepcopy(layout[0]._element))
    for ph in slide.placeholders:
        if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            ph.text_frame.text = str(n)


def deck() -> bytes:
    """Six slides: a title slide; four with a footer drawn as a text box and a
    logo on each, the third of them with a title placeholder under a line;
    and a slide without text."""
    prs = Presentation()
    blank, title_only = prs.slide_layouts[6], prs.slide_layouts[5]

    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Quarterly garden report"
    s.placeholders[1].text = "Prepared by the allotment committee"
    s.notes_slide.notes_text_frame.text = "Welcome everyone.\nIntroduce the committee."

    s = prs.slides.add_slide(blank)
    text_box(s, "Compost results", 0.3, 32)
    text_box(s, "Three bins were turned weekly.", 1.5)
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(2.5), Inches(4), Inches(3),
                       _chart_data()).chart.has_title = False
    text_box(s, FOOTER, 6.8, 10)
    s.shapes.add_picture(png("green"), Inches(9), Inches(6.8))
    slide_number(s, 2)

    s = prs.slides.add_slide(blank)
    text_box(s, "R&D_plan for <next> year", 0.3, 32)
    s.shapes.add_picture(png("red", (60, 60)), Inches(1), Inches(2))
    text_box(s, FOOTER, 6.8, 10)
    s.shapes.add_picture(png("green"), Inches(9), Inches(6.8))
    s.notes_slide.notes_text_frame.text = "Mention the new_shed budget."

    s = prs.slides.add_slide(title_only)
    text_box(s, "A note above the title", 0.1, 12)
    s.shapes.title.text = "Watering rota"
    text_box(s, "Mondays and Thursdays.", 3)
    text_box(s, FOOTER, 6.8, 10)
    s.shapes.add_picture(png("green"), Inches(9), Inches(6.8))

    s = prs.slides.add_slide(blank)
    text_box(s, LONG, 0.3)
    text_box(s, FOOTER, 6.8, 10)
    s.shapes.add_picture(png("green"), Inches(9), Inches(6.8))

    prs.slides.add_slide(blank)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _chart_data():
    data = CategoryChartData()
    data.categories = ["April", "May"]
    data.add_series("Bins", (3, 5))
    return data


def int_deck() -> bytes:
    """One slide with a picture of 32-bit integers (a TIFF in mode I)."""
    pic = Image.new("I", (40, 20))
    pic.putdata([i * 1000 for i in range(40 * 20)])
    buf = io.BytesIO()
    pic.save(buf, format="TIFF")
    buf.seek(0)
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(buf, Inches(1), Inches(1), Inches(4), Inches(2))
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def pptx_converter(*_args, **_kwargs):
    from docling.datamodel.base_models import InputFormat
    from docling.document_converter import DocumentConverter

    return DocumentConverter(allowed_formats=[InputFormat.PPTX])


class SlidesTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        server.converter = pptx_converter  # no PDF models in the tests
        cls.deck = deck()

    def convert(self, **settings):
        args = {"notes": True, "slide_label": "Slide", "notes_label": "Notes"} | settings
        result, files = server.convert_sync("garden.pptx", self.deck, 1 << 24, False, True, 200, 2.0, "pypdfium", args)
        return result, files

    def sections(self, md):
        return re.findall(r"^## (.*)$", md, flags=re.M)

    def test_a_picture_of_32_bit_integers(self):
        # Pillow 13 writes no mode I PNG: the picture comes back as I;16.
        result, files = server.convert_sync("ints.pptx", int_deck(), 1 << 24, False, True, 200, 2.0, "pypdfium", {"notes": False, "slide_label": "Slide", "notes_label": "Notes"})
        self.assertEqual(len(files), 1, result["markdown"])
        self.assertEqual(Image.open(io.BytesIO(files[0][1])).mode, "I;16")

    def test_a_section_per_slide(self):
        result, _ = self.convert()
        self.assertEqual(self.sections(result["markdown"]), [
            "Slide 1 — Quarterly garden report",
            "Slide 2 — Compost results",
            "Slide 3 — R&amp;D\\_plan for &lt;next&gt; year",
            "Slide 4 — Watering rota",
            "Slide 5",
            "Slide 6",
        ])
        self.assertEqual(result["pages"], 6)
        self.assertEqual(result["title"], "Quarterly garden report")

    def test_the_heading_is_not_repeated(self):
        md = self.convert()[0]["markdown"]
        self.assertEqual(md.count("Compost results"), 1)
        self.assertEqual(md.count("Watering rota"), 1)
        self.assertIn("Prepared by the allotment committee", md)
        self.assertIn("A note above the title", md)

    def test_a_long_first_line_stays_in_the_slide(self):
        md = self.convert()[0]["markdown"]
        self.assertIn(" ".join(LONG.split()), md.split("## Slide 5")[1])

    def test_footers_and_logos_left_out(self):
        result, files = self.convert()
        md = result["markdown"]
        self.assertNotIn(FOOTER, md)
        self.assertNotRegex(md.split("## Slide 2")[1].split("## Slide 3")[0], r"(?m)^2$")
        # The logo on four slides of six is gone; the slide's own picture stays.
        self.assertEqual([(p["name"], p["page"]) for p in result["pictures"]], [("fig-1.png", 3)])
        self.assertEqual(len(files), 1)
        self.assertIn("![](fig-1.png)", md.split("## Slide 3")[1].split("## Slide 4")[0])

    def test_the_chart_keeps_its_data(self):
        md = self.convert()[0]["markdown"]
        slide = md.split("## Slide 2")[1].split("## Slide 3")[0]
        self.assertRegex(slide, r"\| April +\| +3 \|")
        self.assertNotIn("fig-", slide)

    def test_notes(self):
        md = self.convert()[0]["markdown"]
        self.assertIn("> **Notes**: Welcome everyone.\n>\n> Introduce the committee.", md)
        self.assertIn("> **Notes**: Mention the new\\_shed budget.", md)
        self.assertEqual(md.count("**Notes**"), 2)

    def test_no_notes(self):
        md = self.convert(notes=False)[0]["markdown"]
        self.assertNotIn("Welcome everyone", md)
        self.assertNotIn(">", md.replace("&gt;", ""))

    def test_labels(self):
        md = self.convert(slide_label="Diapositiva", notes_label="Note del relatore")[0]["markdown"]
        self.assertTrue(md.startswith("## Diapositiva 1 — Quarterly garden report\n"))
        self.assertIn("> **Note del relatore**: Welcome everyone.", md)
        md = self.convert(slide_label="", notes_label="")[0]["markdown"]
        self.assertTrue(md.startswith("## 1 — Quarterly garden report\n"))
        self.assertIn("> Welcome everyone.", md)

    def call(self, **settings):
        args = {"file": {"name": "garden.pptx", "data": base64.b64encode(self.deck).decode()}} | settings
        return asyncio.run(server.call_tool("convert", args))

    def test_the_tool_reads_slides_by_default(self):
        r = self.call()
        self.assertFalse(r.is_error, r.content)
        self.assertIn("## Slide 2 — Compost results", r.structured_content["markdown"])
        self.assertIn("**Notes**", r.structured_content["markdown"])

    def test_the_tool_without_slides(self):
        r = self.call(slides=False, office="native")
        self.assertFalse(r.is_error, r.content)
        md = r.structured_content["markdown"]
        self.assertNotIn("## Slide", md)
        self.assertNotIn("Welcome everyone", md)
        self.assertIn(FOOTER, md)


if __name__ == "__main__":
    unittest.main()
