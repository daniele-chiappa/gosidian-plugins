"""gosidian-docling: a gosidian plugin that turns documents into markdown.

A plugin is an MCP server (gosidian's docs/plugins.md). This one serves
two tools over Streamable HTTP at /mcp:

- gosidian_manifest: the plugin's manifest, its one node docling.convert;
- convert: a document to markdown with docling, the pictures it holds
  given back as PNG files (embedded resources file:///fig-N.png),
  referenced in the markdown by name. The node takes its file by
  reference (file_refs, gosidian 2.81): {name, mime, size, sha256,
  url}, which it fetches from gosidian while the step runs; a file sent
  the old way ({name, mime, data} in base64) still works.

Documents (doc, docx, odt, rtf) are read as docx by docling's Word reader
(office: auto, the default; LibreOffice turns the older formats into
docx), which keeps their tables, their pictures in place and their
right-to-left text, unless drawings sit over their figures: shapes,
arrows and text boxes over a picture, which the Word reader takes apart.
Such a document goes through LibreOffice to PDF, which keeps the figure
whole, as the page shows it; office: pdf and office: native force one way,
and the result says which was taken. Spreadsheets go to xlsx, which
docling reads as tables. Presentations (pptx, and ppt and odp turned into
pptx) are read by docling's pptx reader and written slide by slide: a
section per slide, titled, with its speaker notes, which a PDF loses;
slides: false sends them through PDF, or as pptx with office: native. A
PDF is read with pypdfium2 by default: docling-parse reads twice the
letters of a PDF that draws its bold as two strokes (pdf_backend chooses).
The text docling finds inside a picture of a page (a table of pictograms
with their labels, a diagram's words), which its markdown leaves out,
comes right under the picture (picture_text).

The plugin reads no vault: gosidian hands it the file, or where to fetch
it, and stores the pictures; it calls nothing but that address. Settings
come from the environment:

GOSIDIAN_PLUGIN_TOKEN  the bearer token gosidian sends ("" for none)
DOCLING_MAX_BYTES      the largest document sent in base64, default 20 MiB
DOCLING_MAX_REF_BYTES  the largest document fetched by reference, default 512 MiB
DOCLING_MAX_PAGES      the most pages converted, default 200
DOCLING_WORKERS        conversions at once, default 1
DOCLING_OFFICE_TIMEOUT seconds LibreOffice may take on a document, default 120
DOCLING_OFFICE_MEMORY_MB the memory LibreOffice may take on a document, default 3072 (0: no limit)
PORT                   default 8000
"""

from __future__ import annotations

import base64
import binascii
import collections
import contextlib
import hashlib
import hmac
import html
import io
import json
import logging
import os
import pathlib
import re
import shutil
import subprocess
import tempfile
import threading
import unicodedata
import zipfile

import anyio
import httpx
import jsonschema
import mcp.types as types
import uvicorn
from mcp.server.lowlevel import Server
from mcp.server.streamable_http_manager import StreamableHTTPSessionManager
from mcp.server.transport_security import TransportSecuritySettings
from starlette.applications import Starlette
from starlette.responses import JSONResponse, PlainTextResponse
from starlette.routing import Route

VERSION = os.environ.get("PLUGIN_VERSION", "0.5.3")
TOKEN = os.environ.get("GOSIDIAN_PLUGIN_TOKEN", "")
MAX_BYTES = int(os.environ.get("DOCLING_MAX_BYTES", str(20 << 20)))
MAX_REF_BYTES = int(os.environ.get("DOCLING_MAX_REF_BYTES", str(512 << 20)))
MAX_PAGES = int(os.environ.get("DOCLING_MAX_PAGES", "200"))
WORKERS = max(1, int(os.environ.get("DOCLING_WORKERS", "1")))
OFFICE_TIMEOUT = int(os.environ.get("DOCLING_OFFICE_TIMEOUT", "120"))
OFFICE_MEMORY_MB = int(os.environ.get("DOCLING_OFFICE_MEMORY_MB", "3072"))
PORT = int(os.environ.get("PORT", "8000"))
SOFFICE = shutil.which("soffice")
PRLIMIT = shutil.which("prlimit")

log = logging.getLogger("gosidian-docling")

# The kinds docling reads, by extension.
KINDS = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "html": "text/html",
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
}

# The kinds LibreOffice turns into one docling reads, when it is there: to
# the format of 2007, or to PDF, as office says (a spreadsheet always to
# xlsx, which docling reads as tables).
OFFICE = {
    "doc": ("docx", "application/msword"),
    "rtf": ("docx", "application/rtf"),
    "odt": ("docx", "application/vnd.oasis.opendocument.text"),
    "xls": ("xlsx", "application/vnd.ms-excel"),
    "ods": ("xlsx", "application/vnd.oasis.opendocument.spreadsheet"),
    "ppt": ("pptx", "application/vnd.ms-powerpoint"),
    "odp": ("pptx", "application/vnd.oasis.opendocument.presentation"),
}
if SOFFICE:
    KINDS.update({ext: mime for ext, (_, mime) in OFFICE.items()})

# The documents office concerns: the Word ones, read as docx or through
# PDF (office: auto looks at their drawings), and the presentations with
# slides: false, through PDF or as pptx.
WORD = {"doc", "docx", "odt", "rtf"}
OFFICE_MODES = ("auto", "pdf", "native")

# office: auto. A figure is a picture of at least FIGURE_MIN_CM on its
# longer side (an icon in a table is smaller). A Word document goes through
# PDF when drawings sit over at least DRAWN_SHARE of its figures; a manual
# with numbered callouts on a few figures stays native. The document's
# XML is not read past DOCX_XML_MAX bytes.
FIGURE_MIN_CM = 2.5
DRAWN_SHARE = 1 / 3
DOCX_XML_MAX = 64 << 20

# The text inside a picture of a page comes out under it when it holds at
# least this many words of two letters or more: a table of pictograms, a
# labelled diagram, not a photo with a stray letter or a logo.
PICTURE_TEXT_MIN_WORDS = 4

# The presentations, read slide by slide as pptx (slides, by default).
SLIDES = {"pptx", "ppt", "odp"}
# A slide's first line is its title when it has no title placeholder,
# unless it is longer than this: then the line stays in the slide.
SLIDE_TITLE_MAX = 150

# How a PDF's text is read: pypdfium2, or docling-parse.
PDF_BACKENDS = ("pypdfium", "docling")

PLACEHOLDER = "<!-- gosidian:picture -->"

CONVERT_SETTINGS = {
    "type": "object",
    "required": ["file"],
    "properties": {
        "file": {"type": "string", "description": "The document, a vault path: an attachment of the project, or a file of its IN folder."},
        "ocr": {"type": "boolean", "default": True, "description": "Read the text of scanned pages and pictures."},
        "pictures": {"type": "boolean", "default": True, "description": "Give back the pictures as PNG files."},
        "max_pages": {"type": "integer", "description": f"Convert the first pages only (at most {MAX_PAGES})."},
        "images_scale": {"type": "number", "default": 2.0, "description": "The pictures' scale, 1 = 72 dpi."},
        "pdf_backend": {"type": "string", "enum": list(PDF_BACKENDS), "default": "pypdfium",
                        "description": "How a PDF's text is read: pypdfium (pypdfium2), or docling (docling-parse)."},
        "office": {"type": "string", "enum": list(OFFICE_MODES), "default": "auto",
                   "description": "Documents (doc, docx, odt, rtf): read as docx (native), through LibreOffice to PDF (pdf), or as docx unless drawings sit over their figures, then through PDF (auto). Presentations with slides off: through PDF, or as pptx with native."},
        "picture_text": {"type": "boolean", "default": True,
                         "description": "The text docling finds inside a picture of a page (pictograms and their labels, a diagram's words) right under it, a line for each block, in reading order."},
        "slides": {"type": "boolean", "default": True,
                   "description": "Presentations (pptx, ppt, odp) slide by slide: a section per slide, titled by its title or its first line, with its speaker notes; footers and pictures on more than half the slides left out. Off: as office says."},
        "notes": {"type": "boolean", "default": True, "description": "With slides: a slide's speaker notes under it."},
        "slide_label": {"type": "string", "default": "Slide", "description": "With slides: the word before a slide's number in its heading."},
        "notes_label": {"type": "string", "default": "Notes", "description": "With slides: the word before a slide's speaker notes."},
    },
}

CONVERT_OUTPUT = {
    "type": "object",
    "properties": {
        "markdown": {"type": "string", "description": "The document; a picture is ![caption](fig-N.png)."},
        "title": {"type": "string"},
        "pages": {"type": "integer"},
        "tables": {"type": "integer"},
        "pictures": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {"name": {"type": "string"}, "page": {"type": "integer"}, "caption": {"type": "string"}},
            },
        },
        "office_used": {"type": "string", "enum": ["native", "pdf"],
                        "description": "A document that office concerns: read as docx or pptx (native), or through PDF (pdf)."},
        "office_reason": {"type": "string", "description": "Why it went that way."},
    },
}

MANIFEST = {
    "name": "docling",
    "version": VERSION,
    "gosidian": 1,
    "description": "Documents (" + ", ".join(sorted(k.upper() for k in KINDS if k != "jpeg")) + ") to markdown with docling, with the pictures they hold as PNG files.",
    "nodes": [
        {
            "name": "convert",
            "tool": "convert",
            "description": "A document to markdown, a presentation slide by slide with its speaker notes; its pictures come back as files (output files), named fig-N.png in the markdown.",
            "files": ["file"],
            "file_refs": True,
            "input_schema": CONVERT_SETTINGS,
            "output_schema": CONVERT_OUTPUT,
        }
    ],
    # It fetches the files it is handed from gosidian, nothing else.
    "permissions": {"files": True, "network": ["gosidian"]},
}

# The convert tool as gosidian calls it: file says where to fetch the
# document (url), or carries it (data).
CONVERT_TOOL_INPUT = {
    "type": "object",
    "required": ["file"],
    "properties": {
        "file": {
            "type": "object",
            "required": ["name"],
            "properties": {
                "name": {"type": "string"},
                "mime": {"type": "string"},
                "url": {"type": "string", "description": "Where to fetch it, while the step runs"},
                "size": {"type": "integer"},
                "sha256": {"type": "string"},
                "data": {"type": "string", "description": "base64, the way before file_refs"},
            },
        },
        "ocr": {"type": "boolean"},
        "pictures": {"type": "boolean"},
        "max_pages": {"type": "integer", "minimum": 1},
        "images_scale": {"type": "number", "minimum": 0.5, "maximum": 4},
        "pdf_backend": {"type": "string", "enum": list(PDF_BACKENDS)},
        "office": {"type": "string", "enum": list(OFFICE_MODES)},
        "picture_text": {"type": "boolean"},
        "slides": {"type": "boolean"},
        "notes": {"type": "boolean"},
        "slide_label": {"type": "string", "maxLength": 40},
        "notes_label": {"type": "string", "maxLength": 40},
    },
}


class ConvertError(Exception):
    """A document the plugin cannot convert, said to the caller."""


_converters: dict[tuple, object] = {}
_converters_lock = threading.Lock()
_slots = anyio.Semaphore(WORKERS)


def converter(ocr: bool, scale: float, backend: str = "pypdfium"):
    """A docling converter for these options, built once: it loads models."""
    from docling.datamodel.base_models import InputFormat
    from docling.datamodel.pipeline_options import PdfPipelineOptions
    from docling.document_converter import DocumentConverter, ImageFormatOption, PdfFormatOption

    key = (ocr, scale, backend)
    with _converters_lock:
        if key not in _converters:
            opts = PdfPipelineOptions()
            opts.do_ocr = ocr
            opts.do_table_structure = True
            opts.generate_picture_images = True
            opts.images_scale = scale
            artifacts = os.environ.get("DOCLING_ARTIFACTS_PATH")
            if artifacts:
                opts.artifacts_path = artifacts
            pdf = {"pipeline_options": opts}
            if backend == "pypdfium":
                from docling.backend.pypdfium2_backend import PyPdfiumDocumentBackend

                pdf["backend"] = PyPdfiumDocumentBackend
            _converters[key] = DocumentConverter(
                format_options={
                    InputFormat.PDF: PdfFormatOption(**pdf),
                    InputFormat.IMAGE: ImageFormatOption(pipeline_options=opts),
                }
            )
        return _converters[key]


def office_convert(path: pathlib.Path, into: str, target: str) -> pathlib.Path:
    """Turns an office document into target (pdf, docx, xlsx, pptx), which
    docling reads, with LibreOffice, in the folder into."""
    out = pathlib.Path(into) / "office"
    out.mkdir(exist_ok=True)
    # A profile of its own: two conversions at once do not share one.
    profile = pathlib.Path(into) / "lo-profile"
    try:
        r = subprocess.run(
            office_command(path, out, target, profile),
            capture_output=True, text=True, timeout=OFFICE_TIMEOUT,
        )
    except subprocess.TimeoutExpired:
        raise ConvertError(f"{path.name}: LibreOffice took more than {OFFICE_TIMEOUT} s") from None
    got = out / (path.stem + "." + target)
    if not got.is_file():
        limit = f", with at most {OFFICE_MEMORY_MB} MB (DOCLING_OFFICE_MEMORY_MB)" if office_limited() else ""
        raise ConvertError(f"{path.name}: LibreOffice could not turn it into {target} (exit {r.returncode}{limit}: {(r.stderr or r.stdout).strip()[-200:]})")
    return got


def office_limited() -> bool:
    return OFFICE_MEMORY_MB > 0 and PRLIMIT is not None


def office_command(path: pathlib.Path, out: pathlib.Path, target: str, profile: pathlib.Path) -> list[str]:
    """LibreOffice's command, under prlimit when it has a memory limit: a
    document it cannot open, which takes memory without end, then fails
    alone, instead of taking the container's memory and every conversion in
    flight with it."""
    cmd = [SOFFICE, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--norestore", "--convert-to", target, "--outdir", str(out), str(path)]
    if office_limited():
        cmd = [PRLIMIT, f"--data={OFFICE_MEMORY_MB << 20}", "--"] + cmd
    return cmd


_NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "wpg": "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup",
    "wpc": "http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "v": "urn:schemas-microsoft-com:vml",
    "w10": "urn:schemas-microsoft-com:office:word",
}


def _tag(name: str) -> str:
    prefix, _, local = name.partition(":")
    return "{%s}%s" % (_NS[prefix], local)


_P, _TXBX, _FALLBACK = _tag("w:p"), _tag("w:txbxContent"), _tag("mc:Fallback")
_PIC, _IMAGEDATA, _WSP, _ANCHOR = _tag("pic:pic"), _tag("v:imagedata"), _tag("wps:wsp"), _tag("wp:anchor")
_GROUPS = {_tag("wpg:wgp"), _tag("wpc:wpc"), _tag("v:group")}
_VML_SHAPES = {_tag("v:" + k) for k in ("shape", "rect", "roundrect", "oval", "line", "polyline", "curve", "arc")}
_EMU_CM = 360000
_VML_CM = {"pt": 2.54 / 72, "in": 2.54, "cm": 1.0, "mm": 0.1, "px": 2.54 / 96, "": 2.54 / 72}


def _vml_cm(style: str, key: str) -> float:
    m = re.search(r"(?:^|;)\s*%s\s*:\s*(-?[\d.]+)\s*(pt|in|cm|mm|px|)\s*(?:;|$)" % key, style or "")
    return float(m.group(1)) * _VML_CM[m.group(2)] if m else 0.0


def _picture_cm(el) -> float:
    """A picture's longer side in centimetres; unknown, as large as a figure."""
    if el.tag == _PIC:
        ext = el.find("pic:spPr/a:xfrm/a:ext", _NS)
        if ext is not None:
            try:
                return max(int(ext.get("cx", 0)), int(ext.get("cy", 0))) / _EMU_CM
            except ValueError:
                pass
        return FIGURE_MIN_CM
    style = el.get("style", "")
    return max(_vml_cm(style, "width"), _vml_cm(style, "height")) or FIGURE_MIN_CM


def _holds_picture(el) -> bool:
    return any(True for _ in el.iter(_PIC)) or any(True for _ in el.iter(_IMAGEDATA))


def _is_drawing(el) -> bool:
    """A shape, a line or a text box: not a picture, nor a frame that holds one."""
    if el.tag in _VML_SHAPES and el.find("v:imagedata", _NS) is not None:
        return False
    return (el.tag == _WSP or el.tag in _VML_SHAPES) and not any(_holds_picture(t) for t in el.iter(_TXBX))


def _paragraph_drawings(p) -> tuple[list[float], bool, bool]:
    """A body paragraph's pictures (their sizes, those of the frames anchored
    there among them), whether a drawing floats over the text there (no
    wrapping, or behind it), and whether a group or a canvas there holds a
    picture and a drawing."""
    sizes = [_picture_cm(el) for el in p.iter() if el.tag == _PIC or (el.tag in _VML_SHAPES and el.find("v:imagedata", _NS) is not None)]
    floating = False
    for anchor in p.iter(_ANCHOR):
        if anchor.find("wp:wrapNone", _NS) is not None or anchor.get("behindDoc") in ("1", "true", "on"):
            floating = floating or any(_is_drawing(el) for el in anchor.iter(_WSP))
    for el in p.iter():
        if el.tag in _VML_SHAPES and _is_drawing(el):
            style = el.get("style", "").replace(" ", "")
            wrap = el.find("w10:wrap", _NS)
            if "position:absolute" in style and (wrap is None or wrap.get("type") in (None, "none")):
                floating = True
    grouped = any(_holds_picture(g) and any(_is_drawing(el) for el in g.iter()) for g in p.iter() if g.tag in _GROUPS)
    return sizes, floating, grouped


def office_drawings(docx) -> tuple[bool, str]:
    """Whether a Word document goes through PDF with office: auto, and why.
    It does when drawings sit over at least a third of its figures: shapes,
    lines and arrows, text boxes, grouped with a picture or on a drawing
    canvas with one, or floating over the text (no wrapping) in the
    paragraph that holds a picture, or in the one before or after it. The
    figures are the pictures of FIGURE_MIN_CM and more, a picture in a frame
    among them. docling's Word reader takes such a figure apart, which the
    PDF keeps whole; a manual with numbered callouts on a few of its figures
    stays native, which keeps its tables, its pictures in place and its
    right-to-left text."""
    from lxml import etree

    # No entities, no network, no tree past libxml2's limits.
    parser = etree.XMLParser(resolve_entities=False, no_network=True, remove_comments=True, remove_pis=True)
    try:
        with zipfile.ZipFile(docx) as z:
            info = z.getinfo("word/document.xml")
            if info.file_size > DOCX_XML_MAX:
                return False, f"its text is over {DOCX_XML_MAX >> 20} MiB, its drawings were not looked at"
            with z.open(info) as f:
                root = etree.parse(f, parser).getroot()
    except Exception as e:  # a docx LibreOffice may still read: through PDF
        log.warning("the drawings of %s could not be read: %s", pathlib.Path(str(docx)).name, e)
        return True, f"its drawings could not be read ({type(e).__name__})"
    # Word writes a drawing twice, in the markup of 2010 and in VML for the
    # readers before it (mc:Fallback): the copy does not count.
    for parent in list(root.iter()):
        for child in [c for c in parent if c.tag == _FALLBACK]:
            parent.remove(child)
    # The body's paragraphs, those of the tables among them; a text box's
    # paragraphs belong to the paragraph that anchors it.
    paragraphs, stack = [], [root]
    while stack:
        el = stack.pop()
        if el.tag == _P:
            paragraphs.append(el)
        else:
            stack.extend(reversed(list(el)))
    found = [_paragraph_drawings(p) for p in paragraphs]
    figures = drawn = 0
    for i, (sizes, _, grouped) in enumerate(found):
        if not (grouped or any(s >= FIGURE_MIN_CM for s in sizes)):
            continue
        figures += 1
        if grouped or any(found[j][1] for j in range(max(0, i - 1), min(len(found), i + 2))):
            drawn += 1
    if not figures:
        return False, "no figures"
    if not drawn:
        return False, f"no drawings over its {figures} figures" if figures > 1 else "no drawings over its figure"
    said = f"drawings over {drawn} of its {figures} figures"
    if drawn >= figures * DRAWN_SHARE:
        return True, said
    return False, said + ", under a third"


def office_source(path: pathlib.Path, ext: str, office: str, into: str) -> tuple[pathlib.Path, str, str]:
    """The file docling reads for a Word document, or a presentation read
    without slides, as office says (auto, pdf, native), made by LibreOffice
    in the folder into; and the way it took, native or pdf, and why."""
    native = OFFICE[ext][0] if ext in OFFICE else None  # docx or pptx
    if office == "native" or not SOFFICE:
        reason = "office: native" if office == "native" else "no LibreOffice to make a PDF"
        return (office_convert(path, into, native) if native else path), "native", reason
    if office == "pdf":
        return office_convert(path, into, "pdf"), "pdf", "office: pdf"
    if ext not in WORD:
        # A slide drawn with shapes and arrows stays whole through PDF.
        return office_convert(path, into, "pdf"), "pdf", "a presentation without slides goes through PDF"
    docx = office_convert(path, into, "docx") if native else path
    pdf, reason = office_drawings(docx)
    if pdf:
        return office_convert(path, into, "pdf"), "pdf", reason
    return docx, "native", reason


def one_line(text: str) -> str:
    return " ".join((text or "").split())


def md_text(text: str) -> str:
    """A slide's title or a line of its notes in markdown, escaped as docling
    escapes the text it writes."""
    return html.escape(re.sub(r"(?<!\\)_", r"\_", one_line(text)), quote=False)


def slide_placeholders(source) -> dict[int, set[str]]:
    """The lines of each slide's footer, date and number placeholders, by
    slide: docling's pptx reader gives them as lines of the slide."""
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER

    kinds = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.SLIDE_NUMBER}
    found: dict[int, set[str]] = {}
    try:
        deck = Presentation(io.BytesIO(source) if isinstance(source, bytes) else str(source))
        for n, slide in enumerate(deck.slides, 1):
            for shape in slide.placeholders:
                if shape.placeholder_format.type in kinds and shape.has_text_frame:
                    found.setdefault(n, set()).update(one_line(p.text) for p in shape.text_frame.paragraphs)
    except Exception as e:  # docling read it: the slides come out, with these lines
        log.warning("the placeholders of the slides could not be read: %s", e)
    return found


def slides_markdown(doc, source, notes: bool, slide_label: str, notes_label: str) -> tuple[str, str]:
    """A presentation read by docling's pptx reader, as markdown, a section
    per slide: its heading, with the slide's title (its title placeholder,
    or else its first line), what the slide holds, and its speaker notes.
    The footer, date and number placeholders, and the lines and pictures on
    more than half the slides (a footer drawn as a text box, a logo), go to
    the furniture, which the markdown leaves out; docling reads neither the
    layouts nor the masters, nor their pictures. Gives the markdown, its
    pictures as PLACEHOLDER, and the first title."""
    from docling_core.types.doc import ContentLayer, DocItemLabel, ImageRefMode, PictureItem, TextItem

    def repeated(item):
        """What a footer or a logo repeats from slide to slide: its line, or its picture."""
        if isinstance(item, TextItem) and item.label not in (DocItemLabel.TITLE, DocItemLabel.CAPTION):
            return ("line", one_line(item.text)) if one_line(item.text) else None
        if isinstance(item, PictureItem) and item.image is not None:
            return ("picture", hashlib.sha256(str(item.image.uri).encode()).hexdigest())
        return None

    numbers = sorted(doc.pages)
    on = {n: [] for n in numbers}
    for item, _level in doc.iterate_items():
        if item.prov and item.prov[0].page_no in on:
            on[item.prov[0].page_no].append(item)
    seen = collections.Counter(k for n in numbers for k in {repeated(i) for i in on[n]} if k)
    common = {k for k, c in seen.items() if c * 2 > len(numbers)} if len(numbers) >= 3 else set()
    placeholders = slide_placeholders(source)

    titles = {}
    for n in numbers:
        for item in on[n]:
            if repeated(item) in common or (isinstance(item, TextItem) and one_line(item.text) in placeholders.get(n, ())):
                item.content_layer = ContentLayer.FURNITURE
        lines = [i for i in on[n] if isinstance(i, TextItem) and i.label != DocItemLabel.CAPTION and i.content_layer == ContentLayer.BODY and one_line(i.text)]
        head = [i for i in lines if i.label == DocItemLabel.TITLE]
        if not head and lines and len(one_line(lines[0].text)) <= SLIDE_TITLE_MAX:
            head = lines[:1]
        for i in head:
            i.content_layer = ContentLayer.FURNITURE  # it is the heading
        titles[n] = " ".join(one_line(i.text) for i in head)

    said = collections.defaultdict(list)
    if notes:
        for item, _level in doc.iterate_items(included_content_layers={ContentLayer.NOTES}):
            if isinstance(item, TextItem) and item.label == DocItemLabel.TEXT and item.prov:
                said[item.prov[0].page_no].extend(item.text.replace("\v", "\n").split("\n"))

    out = []
    for n in numbers:
        heading = " ".join(x for x in (slide_label.strip(), str(n)) if x)
        out.append(f"## {heading} — {md_text(titles[n])}" if titles[n] else f"## {heading}")
        body = doc.export_to_markdown(page_no=n, image_mode=ImageRefMode.PLACEHOLDER, image_placeholder=PLACEHOLDER).strip()
        if body:
            out.append(body)
        lines = [md_text(x) for x in said[n] if x.strip()]
        if lines:
            if notes_label.strip():
                lines[0] = f"**{md_text(notes_label)}**: {lines[0]}"
            out.append("\n>\n".join("> " + x for x in lines))
    return "\n\n".join(out), next((t for t in titles.values() if t), "")


def md_line(text: str) -> str:
    """A line of a picture's text in markdown: escaped as docling escapes
    its text, and not read as a heading, a list item or a rule."""
    s = md_text(text)
    m = re.match(r"(#{1,6}|[-+*]|\d{1,9}[.)])(?=\s|$)", s)
    if m:
        mark = m.group(1)
        s = (mark[:-1] + "\\" + mark[-1] if mark[0].isdigit() else "\\" + mark) + s[len(mark):]
    elif re.fullmatch(r"[-=_*](?: *[-=_*]){2,}", s):
        s = "\\" + s
    return s


def _xy_gaps(boxes: list, lo: int, hi: int) -> tuple[list, list[float]]:
    """The boxes sorted along an axis (lo, hi: their sides), and the gap
    before each one after the first, from the furthest end before it
    (negative where they overlap)."""
    boxes = sorted(boxes, key=lambda b: b[lo])
    gaps, end = [], boxes[0][hi]
    for b in boxes[1:]:
        gaps.append(b[lo] - end)
        end = max(end, b[hi])
    return boxes, gaps


def _split(boxes: list, gaps: list[float], least: float) -> list[list]:
    parts = [[boxes[0]]]
    for b, gap in zip(boxes[1:], gaps):
        if gap >= least:
            parts.append([])
        parts[-1].append(b)
    return parts


def picture_lines(item, doc) -> list[str]:
    """The text docling found inside a picture of a page, which its markdown
    leaves out: the text layer's or the OCR's lines that fall in the
    picture, which docling keeps as its children. Pieces side by side make
    a line; lines one under the other, close and aligned (a label on two
    lines), a block; the blocks come in reading order, a line each, cut
    by the widest gaps between them: two columns of labels one after the
    other, a row of boxes left to right (right to left in Arabic and
    Hebrew). A block without letters is left out. Nothing when the picture
    holds fewer than PICTURE_TEXT_MIN_WORDS words of two letters or more,
    or is a chart, most of its blocks numbers."""
    skip = {r.cref for r in (*item.captions, *item.footnotes)}
    boxes, todo = [], list(item.children)
    while todo:
        ref = todo.pop()
        if ref.cref in skip:
            continue
        child = ref.resolve(doc)
        todo.extend(child.children)
        text = one_line(getattr(child, "text", "") or "")
        if not text or not getattr(child, "prov", None):
            continue
        prov = child.prov[0]
        page = doc.pages.get(prov.page_no)
        bb = prov.bbox.to_top_left_origin(page.size.height) if page is not None else prov.bbox
        boxes.append((bb.l, min(bb.t, bb.b), bb.r, max(bb.t, bb.b), text))
    if len(re.findall(r"[^\W\d_]{2,}", " ".join(b[4] for b in boxes))) < PICTURE_TEXT_MIN_WORDS:
        return []

    def height(b):
        return max(b[3] - b[1], 1.0)

    # Right-to-left writing (Arabic, Hebrew): a line's pieces and the
    # columns go from the right.
    letters = [unicodedata.bidirectional(c) for b in boxes for c in b[4] if c.isalpha()]
    rtl = sum(k in ("R", "AL") for k in letters) * 2 > len(letters)

    def same_row(b, row) -> bool:
        h, small = max(height(b), height(row)), min(height(b), height(row))
        overlap = min(b[3], row[3]) - max(b[1], row[1])
        gap = max(b[0] - row[2], row[0] - b[2])
        if gap < -0.5 * small:
            return False
        # A mark beside a word (an apostrophe, a degree sign), or words of a
        # like size, however justified the line, not the next column's.
        return (overlap >= 0.3 * small and gap <= 0.5 * h) or (overlap >= 0.5 * small and gap <= 3 * h and small >= 0.75 * h)

    # Pieces on the same line: a line.
    rows: list[list] = []
    for b in sorted(boxes, key=lambda b: (b[1], b[0])):
        for row in rows:
            if same_row(b, row[0]):
                row[0] = (min(b[0], row[0][0]), min(b[1], row[0][1]), max(b[2], row[0][2]), max(b[3], row[0][3]), "")
                row.append(b)
                break
        else:
            rows.append([(b[0], b[1], b[2], b[3], ""), b])
    lines = []
    for row in rows:
        pieces = sorted(row[1:], key=lambda b: -b[2] if rtl else b[0])
        text = pieces[0][4]
        for prev, b in zip(pieces, pieces[1:]):
            gap = prev[0] - b[2] if rtl else b[0] - prev[2]
            text += (" " if gap >= 0.3 * max(height(prev), height(b)) else "") + b[4]
        lines.append((*row[0][:4], text))

    # Lines one under the other, close and aligned: a block.
    blocks: list[list] = []
    for line in sorted(lines, key=lambda b: (b[1], b[0])):
        best = None
        for block in blocks:
            last = block[-1]
            h = max(height(line), height(last))
            gap = line[1] - last[3]
            overlap = min(line[2], last[2]) - max(line[0], last[0])
            if (-0.25 * h <= gap <= h and 2 / 3 <= height(line) / height(last) <= 1.5
                    and (overlap >= 0.3 * min(line[2] - line[0], last[2] - last[0]) or abs(line[0] - last[0]) <= h)
                    and (best is None or gap < best[0])):
                best = (gap, block)
        if best:
            best[1].append(line)
        else:
            blocks.append([line])
    merged = []
    for block in blocks:
        text = block[0][4]
        for line in block[1:]:
            if len(text) > 1 and text.endswith("-") and text[-2].isalpha() and line[4][:1].islower():
                text = text[:-1] + line[4]  # a word split at the line's end
            else:
                text += " " + line[4]
        merged.append((min(b[0] for b in block), block[0][1], max(b[2] for b in block), block[-1][3], text))
    # A block without letters (a bar's value, an axis' tick, a callout's
    # number) says nothing without the picture; a picture made mostly of
    # them is a chart, which keeps none.
    worded = [b for b in merged if any(c.isalpha() for c in b[4])]
    if len(worded) * 2 < len(merged):
        return []
    merged = worded

    # Reading order: cut at the widest gap, across (between columns, at
    # least a line's height) or down (between rows), and again in each
    # part. A gap across weighs twice: the columns of a page the picture
    # swallowed are read one after the other, not row by row.
    unit = sorted(height(b) for b in lines)[len(lines) // 2]
    order, stack = [], [merged]
    while stack:
        group = stack.pop()
        if len(group) == 1:
            order.append(group[0])
            continue
        across, gaps_x = _xy_gaps(group, 0, 2)
        down, gaps_y = _xy_gaps(group, 1, 3)
        wide_x, wide_y = max(gaps_x), max(gaps_y)
        if wide_x >= unit and 2 * wide_x >= wide_y:
            parts = _split(across, gaps_x, wide_x)
            if rtl:
                parts.reverse()
        elif wide_y > 0:
            parts = _split(down, gaps_y, wide_y)
        else:
            order.extend(sorted(group, key=lambda b: (b[1], -b[2] if rtl else b[0])))
            continue
        stack.extend(reversed(parts))
    return [md_line(b[4]) for b in order]


def convert_sync(name: str, source, size_cap: int, ocr: bool, pictures: bool, max_pages: int, scale: float, backend: str = "pypdfium", slides: dict | None = None, picture_text: bool = True) -> tuple[dict, list[tuple[str, bytes]]]:
    """Converts one document, its bytes or a file named as it: the structured
    result and the picture files. slides, for a pptx, holds slides_markdown's
    notes, slide_label and notes_label. picture_text writes the text inside
    a picture of a page (a PDF's, an image's) under it."""
    from docling.datamodel.base_models import DocumentStream, InputFormat
    from docling_core.types.doc import ImageRefMode, PictureItem, TableItem

    raw = source
    if isinstance(source, bytes):
        source = DocumentStream(name=name, stream=io.BytesIO(source))
    result = converter(ocr, scale, backend).convert(source, max_num_pages=max_pages, max_file_size=size_cap, raises_on_error=False)
    doc = result.document
    if doc is None or (result.status.name not in ("SUCCESS", "PARTIAL_SUCCESS")):
        errors = "; ".join(str(e.error_message) for e in (result.errors or [])) or result.status.name
        raise ConvertError(f"docling could not convert {name}: {errors}")

    title = None
    if slides is not None:
        md, title = slides_markdown(doc, raw, **slides)
    else:
        md = doc.export_to_markdown(image_mode=ImageRefMode.PLACEHOLDER, image_placeholder=PLACEHOLDER)
    # Only the pages docling lays out keep the text inside their pictures.
    paged = picture_text and getattr(getattr(result, "input", None), "format", None) in (InputFormat.PDF, InputFormat.IMAGE)
    files: list[tuple[str, bytes]] = []
    shown: list[dict] = []
    tables = 0
    parts = md.split(PLACEHOLDER)
    out = [parts[0]]
    n = 0
    for item, _level in doc.iterate_items():
        if isinstance(item, TableItem):
            tables += 1
        if not isinstance(item, PictureItem):
            continue
        if n + 1 >= len(parts):
            break  # a picture the markdown does not show
        n += 1
        caption = item.caption_text(doc) or ""
        ref = ""
        done = "".join(out)
        in_cell = done[done.rfind("\n") + 1 :].lstrip().startswith("|")
        img = item.get_image(doc) if pictures else None
        if img is not None:
            fname = f"fig-{len(files) + 1}.png"
            buf = io.BytesIO()
            img.save(buf, format="PNG")
            files.append((fname, buf.getvalue()))
            page = item.prov[0].page_no if item.prov else None
            shown.append({"name": fname, "page": page, "caption": caption})
            alt = caption.replace("[", "(").replace("]", ")").replace("\n", " ")
            if in_cell:
                # In a table's cell: on its line, or the table breaks.
                ref = f" ![{alt}]({fname}) "
            else:
                # A picture is a block of its own, not the tail of a list item.
                ref = f"\n\n![{alt}]({fname})\n\n"
        if paged and not in_cell:
            # Its text right under it, which the markdown left out.
            lines = picture_lines(item, doc)
            if lines:
                ref += "\n\n" + "\n\n".join(lines) + "\n\n"
        out.append(ref)
        out.append(parts[n])
    out.extend(parts[n + 1 :])

    if title is None:
        for item, _level in doc.iterate_items():
            label = getattr(item, "label", None)
            if label is not None and str(getattr(label, "value", label)) in ("title", "section_header"):
                title = getattr(item, "text", None)
                break
    return {
        "markdown": re.sub(r"\n{3,}", "\n\n", "".join(out)).strip() + "\n",
        "title": title or "",
        "pages": len(doc.pages),
        "tables": tables,
        "pictures": shown,
    }, files


def file_name(f) -> str:
    """The document's name, its kind checked."""
    if not isinstance(f, dict):
        raise ConvertError("file is {name, mime, url} (or data): gosidian sends it for a node's file setting")
    name = os.path.basename(str(f.get("name", "")).replace("\\", "/")) or "document"
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if ext not in KINDS:
        raise ConvertError(f"{name}: docling reads {', '.join(sorted(KINDS))}")
    return name


def fetch_file(f: dict, name: str, into: str) -> pathlib.Path:
    """Fetches a document handed by reference into the folder into, under its
    name, checked against the size and the SHA-256 gosidian gave."""
    url = str(f.get("url", ""))
    if not url.startswith(("http://", "https://")):
        raise ConvertError(f"{name}: url {url!r} is not an http(s) address")
    size = int(f.get("size") or 0)
    if size > MAX_REF_BYTES:
        raise ConvertError(f"{name} is {size} bytes, at most {MAX_REF_BYTES}")
    want = str(f.get("sha256", "")).lower()
    path = pathlib.Path(into) / name
    h = hashlib.sha256()
    got = 0
    try:
        with open(path, "wb") as out, httpx.stream("GET", url, timeout=httpx.Timeout(60.0, read=300.0), follow_redirects=False) as r:
            if r.status_code != 200:
                raise ConvertError(f"{name}: fetching it answered {r.status_code}")
            for chunk in r.iter_bytes(1 << 20):
                got += len(chunk)
                if got > MAX_REF_BYTES:
                    raise ConvertError(f"{name} is more than {MAX_REF_BYTES} bytes")
                h.update(chunk)
                out.write(chunk)
    except httpx.HTTPError as e:
        raise ConvertError(f"{name}: fetching it failed ({e})") from None
    if not got:
        raise ConvertError(f"{name} is empty")
    if want and h.hexdigest() != want:
        raise ConvertError(f"{name}: what came is not the file gosidian named (its SHA-256 differs)")
    return path


def decode_file(args: dict) -> tuple[str, bytes]:
    f = args.get("file")
    name = file_name(f)
    try:
        data = base64.b64decode(str(f.get("data", "")), validate=True)
    except (binascii.Error, ValueError) as e:
        raise ConvertError(f"{name}: data is not base64 ({e})") from None
    if not data:
        raise ConvertError(f"{name} is empty")
    if len(data) > MAX_BYTES:
        raise ConvertError(f"{name} is {len(data)} bytes, at most {MAX_BYTES}")
    return name, data


TOOLS = [
    types.Tool(name="gosidian_manifest", description="The manifest of this gosidian plugin.", input_schema={"type": "object"}),
    types.Tool(
        name="convert",
        description="A document to markdown with docling; its pictures come back as embedded PNG resources (file:///fig-N.png).",
        input_schema=CONVERT_TOOL_INPUT,
    ),
]


async def on_list_tools(_ctx, _params) -> types.ListToolsResult:
    return types.ListToolsResult(tools=TOOLS)


def tool_error(text: str) -> types.CallToolResult:
    return types.CallToolResult(content=[types.TextContent(type="text", text=text)], is_error=True)


async def on_call_tool(_ctx, params: types.CallToolRequestParams) -> types.CallToolResult:
    """A tools/call answered as the SDK's 1.x decorator did, which 2.x
    leaves to the server: the arguments checked against the tool's schema,
    the manifest as JSON text and structured content, and a failure as a
    tool error with its message (2.x turns an exception into a bare
    "Internal server error")."""
    arguments = params.arguments or {}
    tool = next((t for t in TOOLS if t.name == params.name), None)
    if tool is not None:
        try:
            jsonschema.validate(instance=arguments, schema=tool.input_schema)
        except jsonschema.ValidationError as e:
            return tool_error(f"Input validation error: {e.message}")
    try:
        result = await call_tool(params.name, arguments)
    except Exception as e:
        return tool_error(str(e))
    if isinstance(result, dict):
        return types.CallToolResult(content=[types.TextContent(type="text", text=json.dumps(result, indent=2))], structured_content=result)
    return result


async def call_tool(name: str, arguments: dict):
    if name == "gosidian_manifest":
        return MANIFEST
    if name != "convert":
        raise ValueError(f"no tool {name}")
    tmp = None
    try:
        f = arguments.get("file")
        if isinstance(f, dict) and f.get("url"):
            fname = file_name(f)
            tmp = tempfile.mkdtemp(prefix="docling-")
            source = await anyio.to_thread.run_sync(fetch_file, f, fname, tmp)
            cap = MAX_REF_BYTES
        else:
            fname, source = decode_file(arguments)
            cap = MAX_BYTES
        max_pages = min(int(arguments.get("max_pages") or MAX_PAGES), MAX_PAGES)
        ocr = bool(arguments.get("ocr", True))
        pictures = bool(arguments.get("pictures", True))
        scale = float(arguments.get("images_scale") or 2.0)
        backend = str(arguments.get("pdf_backend") or "pypdfium")
        if backend not in PDF_BACKENDS:
            raise ConvertError(f"pdf_backend {backend!r}: {' or '.join(PDF_BACKENDS)}")
        picture_text = arguments.get("picture_text") is not False
        async with _slots:
            ext = fname.rsplit(".", 1)[-1].lower()
            office = str(arguments.get("office") or "auto")
            if office not in OFFICE_MODES:
                raise ConvertError(f"office {office!r}: {', '.join(OFFICE_MODES)}")
            # A presentation goes slide by slide, as pptx: a PDF loses its notes.
            slides = None
            if ext in SLIDES and arguments.get("slides") is not False:
                slides = {"notes": arguments.get("notes") is not False}
                for key, default in (("slide_label", "Slide"), ("notes_label", "Notes")):
                    got = arguments.get(key)
                    slides[key] = got[:40] if isinstance(got, str) else default
            way = None
            if ext in WORD or (ext in SLIDES and slides is None) or ext in OFFICE:
                if tmp is None:
                    tmp = tempfile.mkdtemp(prefix="docling-")
                if isinstance(source, bytes):
                    path = pathlib.Path(tmp) / fname
                    path.write_bytes(source)
                    source = path
                if ext in WORD or (ext in SLIDES and slides is None):
                    source, *way = await anyio.to_thread.run_sync(office_source, pathlib.Path(source), ext, office, tmp)
                else:  # a spreadsheet, or a presentation slide by slide
                    source = await anyio.to_thread.run_sync(office_convert, pathlib.Path(source), tmp, OFFICE[ext][0])
            result, files = await anyio.to_thread.run_sync(convert_sync, fname, source, cap, ocr, pictures, max_pages, scale, backend, slides, picture_text)
            if way:
                result["office_used"], result["office_reason"] = way
                log.info("%s went %s: %s", fname, *way)
    except ConvertError as e:
        return tool_error(str(e))
    finally:
        if tmp is not None:
            shutil.rmtree(tmp, ignore_errors=True)
    log.info("converted %s: %d pages, %d pictures", fname, result["pages"], len(files))
    content: list[types.ContentBlock] = [types.TextContent(type="text", text=result["markdown"])]
    for fn, blob in files:
        content.append(
            types.EmbeddedResource(
                type="resource",
                resource=types.BlobResourceContents(uri=f"file:///{fn}", mime_type="image/png", blob=base64.b64encode(blob).decode()),
            )
        )
    return types.CallToolResult(content=content, structured_content=result)


app = Server("gosidian-docling", version=VERSION, on_list_tools=on_list_tools, on_call_tool=on_call_tool)

sessions = StreamableHTTPSessionManager(
    app=app,
    json_response=True,
    stateless=True,
    # The bearer token guards the plugin; the Host header differs by network.
    security_settings=TransportSecuritySettings(enable_dns_rebinding_protection=False),
    # A 20 MiB document is 27 MiB in base64.
    max_request_body_size=MAX_BYTES * 4 // 3 + (1 << 20),
)


class MCPEndpoint:
    async def __call__(self, scope, receive, send):
        await sessions.handle_request(scope, receive, send)


async def health(_request):
    return JSONResponse({"status": "ok", "plugin": "docling", "version": VERSION})


class BearerAuth:
    """Wants the token on everything but /health."""

    def __init__(self, inner, token: str):
        self.inner, self.token = inner, token

    async def __call__(self, scope, receive, send):
        if scope["type"] == "http" and self.token and scope.get("path") != "/health":
            got = dict(scope.get("headers") or []).get(b"authorization", b"").decode("latin-1")
            if not hmac.compare_digest(got.encode(), f"Bearer {self.token}".encode()):
                await PlainTextResponse("unauthorized", status_code=401)(scope, receive, send)
                return
        await self.inner(scope, receive, send)


@contextlib.asynccontextmanager
async def lifespan(_app):
    async with sessions.run():
        # Load the models now, not at the first document.
        await anyio.to_thread.run_sync(converter, True, 2.0)
        log.info("gosidian-docling %s ready on :%d (token %s)", VERSION, PORT, "on" if TOKEN else "off")
        if SOFFICE and OFFICE_MEMORY_MB > 0 and not PRLIMIT:
            log.warning("no prlimit: LibreOffice runs without a memory limit")
        yield


asgi = BearerAuth(
    Starlette(routes=[Route("/health", health), Route("/mcp", endpoint=MCPEndpoint())], lifespan=lifespan),
    TOKEN,
)

if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(name)s: %(message)s")
    uvicorn.run(asgi, host="0.0.0.0", port=PORT, log_level="info")
