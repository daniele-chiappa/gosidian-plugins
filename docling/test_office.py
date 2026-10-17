"""Tests of Word documents (office: auto, on docx made here with python-docx
and the XML of their drawings) and of the text inside a picture of a page
(picture_text, on a document built here as docling's PDF pipeline builds
it): python -m unittest test_office (scripts/test.sh runs them). They
need docling's Word reader (docling-slim[convert-core,format-docx]) and the
server's own packages, not the models of the image nor LibreOffice."""

import asyncio
import base64
import copy
import io
import os
import pathlib
import shutil
import sys
import tempfile
import unittest
from types import SimpleNamespace
from unittest import mock

from docx import Document
from docx.oxml import parse_xml
from docx.shared import Cm
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import server  # noqa: E402

NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:wpg="http://schemas.microsoft.com/office/word/2010/wordprocessingGroup" '
    'xmlns:wpc="http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:o="urn:schemas-microsoft-com:office:office" '
    'xmlns:w10="urn:schemas-microsoft-com:office:word" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
)
WPS_URI = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
_ids = iter(range(100, 10000))


def png(color="teal", size=(120, 80)) -> io.BytesIO:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, format="PNG")
    buf.seek(0)
    return buf


def shape(kind="arrow", text=None, wrap="<wp:wrapNone/>", picture_inside=None) -> str:
    """A floating DrawingML shape: an arrow, a rectangle, or a text box
    (text), which may hold a picture (a frame)."""
    n = next(_ids)
    geom = {"arrow": "straightConnector1", "rect": "rect", "callout": "borderCallout1"}[kind]
    line = '<a:ln w="19050"><a:tailEnd type="triangle"/></a:ln>' if kind == "arrow" else ""
    body = ""
    if text is not None or picture_inside is not None:
        inner = f"<w:p><w:r><w:t>{text or ''}</w:t></w:r></w:p>"
        if picture_inside is not None:
            inner = f"<w:p><w:r>{picture_inside}</w:r></w:p>" + inner
        body = f"<wps:txbx><w:txbxContent>{inner}</w:txbxContent></wps:txbx>"
    nv = "<wps:cNvCnPr/>" if kind == "arrow" else '<wps:cNvSpPr txBox="1"/>'
    return (
        f'<w:drawing {NS}><wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="{n}" '
        'behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="column"><wp:posOffset>720000</wp:posOffset></wp:positionH>'
        '<wp:positionV relativeFrom="paragraph"><wp:posOffset>-1440000</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="1080000" cy="360000"/><wp:effectExtent l="0" t="0" r="0" b="0"/>{wrap}'
        f'<wp:docPr id="{n}" name="Shape {n}"/><wp:cNvGraphicFramePr/>'
        f'<a:graphic><a:graphicData uri="{WPS_URI}"><wps:wsp>{nv}'
        f'<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1080000" cy="360000"/></a:xfrm>'
        f'<a:prstGeom prst="{geom}"><a:avLst/></a:prstGeom>{line}</wps:spPr>{body}<wps:bodyPr/></wps:wsp>'
        "</a:graphicData></a:graphic></wp:anchor></w:drawing>"
    )


def picture(doc, cm=6.0):
    """A paragraph with an inline picture of cm centimetres."""
    doc.add_picture(png(), width=Cm(cm))
    return doc.paragraphs[-1]


def pic_xml(doc, cm=6.0) -> str:
    """The pic:pic of a picture, its image a part of doc, as XML."""
    p = picture(doc, cm)
    el = copy.deepcopy(p._p.xpath(".//pic:pic")[0])
    p._p.getparent().remove(p._p)
    from lxml import etree

    return etree.tostring(el).decode()


def drawing_in(paragraph, xml: str):
    paragraph.add_run()._r.append(parse_xml(xml))


def save(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def plain() -> bytes:
    """Two figures and some text, nothing drawn over them."""
    doc = Document()
    doc.add_heading("Seed drill care", 1)
    doc.add_paragraph("Clean the hopper after each sowing.")
    picture(doc)
    doc.add_paragraph("The drive chain wants oil every season.")
    picture(doc, 8)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Hopper"
    table.cell(0, 1).text = "Empty it before winter."
    return save(doc)


def arrows() -> bytes:
    """A figure with an arrow and a text box over it."""
    doc = Document()
    doc.add_paragraph("Fit the coulter as shown.")
    p = picture(doc)
    drawing_in(p, shape("arrow"))
    drawing_in(p, shape("rect", text="Coulter bolt"))
    doc.add_paragraph("Tighten the bolt by hand.")
    return save(doc)


def label_after() -> bytes:
    """A figure whose label floats from the next paragraph."""
    doc = Document()
    picture(doc)
    p = doc.add_paragraph("Turn the wheel.")
    drawing_in(p, shape("callout", text="2"))
    return save(doc)


def grouped(kind="wpg") -> bytes:
    """A picture and an arrow grouped (wpg), or on a drawing canvas (wpc)."""
    doc = Document()
    doc.add_paragraph("The frame, from above.")
    pic = pic_xml(doc)
    arrow = (
        '<wps:wsp><wps:cNvCnPr/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="900000" cy="0"/></a:xfrm>'
        '<a:prstGeom prst="straightConnector1"><a:avLst/></a:prstGeom><a:ln><a:tailEnd type="triangle"/></a:ln></wps:spPr><wps:bodyPr/></wps:wsp>'
    )
    if kind == "wpg":
        uri, inner = "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup", f"<wpg:wgp><wpg:cNvGrpSpPr/><wpg:grpSpPr/>{pic}{arrow}</wpg:wgp>"
    else:
        uri, inner = "http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas", f"<wpc:wpc>{pic}{arrow}</wpc:wpc>"
    xml = (
        f'<w:drawing {NS}><wp:inline distT="0" distB="0" distL="0" distR="0"><wp:extent cx="2160000" cy="1440000"/>'
        f'<wp:docPr id="{next(_ids)}" name="Group"/><a:graphic><a:graphicData uri="{uri}">{inner}</a:graphicData></a:graphic></wp:inline></w:drawing>'
    )
    drawing_in(doc.add_paragraph(), xml)
    return save(doc)


def vml() -> bytes:
    """The markup of before 2010: a VML picture, a VML arrow over it."""
    doc = Document()
    rid, _ = doc.part.get_or_add_image(png())
    p = doc.add_paragraph()
    drawing_in(p, f'<w:pict {NS}><v:shape id="pic1" style="width:170pt;height:113pt"><v:imagedata r:id="{rid}" o:title=""/></v:shape></w:pict>')
    drawing_in(p, f'<w:pict {NS}><v:line id="l1" from="10pt,10pt" to="90pt,40pt" style="position:absolute;z-index:2"><v:stroke endarrow="block"/></v:line></w:pict>')
    doc.add_paragraph("Grease the hub.")
    return save(doc)


def alternate() -> bytes:
    """A shape written twice, as Word does (mc:AlternateContent): a text box
    with a picture (a frame), and its VML copy, a plain text box."""
    doc = Document()
    p = doc.add_paragraph("Before the frame.")
    frame = shape("rect", text="Figure 1", picture_inside=pic_xml(doc).replace("<pic:pic", f"<w:drawing {NS}><wp:inline><wp:extent cx=\"2160000\" cy=\"1440000\"/><wp:docPr id=\"{next(_ids)}\" name=\"P\"/><a:graphic><a:graphicData uri=\"http://schemas.openxmlformats.org/drawingml/2006/picture\"><pic:pic", 1).replace("</pic:pic>", "</pic:pic></a:graphicData></a:graphic></wp:inline></w:drawing>"))
    fallback = '<w:pict><v:rect style="position:absolute;width:85pt;height:28pt"><v:textbox><w:txbxContent><w:p><w:r><w:t>Figure 1</w:t></w:r></w:p></w:txbxContent></v:textbox></v:rect></w:pict>'
    drawing_in(p, f'<mc:AlternateContent {NS}><mc:Choice Requires="wps">{frame}</mc:Choice><mc:Fallback>{fallback}</mc:Fallback></mc:AlternateContent>')
    return save(doc)


def sidebar() -> bytes:
    """A figure next to a text box the text wraps around: not over it."""
    doc = Document()
    p = picture(doc)
    drawing_in(p, shape("rect", text="Tip: oil the chain in spring.", wrap='<wp:wrapSquare wrapText="bothSides"/>'))
    return save(doc)


def manual() -> bytes:
    """Six figures, numbered callouts over one: a manual."""
    doc = Document()
    for n in range(6):
        doc.add_paragraph(f"Step {n + 1}: check the tines.")
        p = picture(doc)
        if n == 2:
            for k in range(3):
                drawing_in(p, shape("callout", text=str(k + 1)))
    return save(doc)


def icons() -> bytes:
    """Small pictures (icons) with a shape near them: no figures."""
    doc = Document()
    p = picture(doc, 1.2)
    drawing_in(p, shape("arrow"))
    return save(doc)


class DrawingsTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp)

    def judge(self, data: bytes):
        path = pathlib.Path(self.tmp) / "doc.docx"
        path.write_bytes(data)
        return server.office_drawings(path)

    def test_nothing_drawn_reads_native(self):
        self.assertEqual(self.judge(plain()), (False, "no drawings over its 2 figures"))

    def test_an_arrow_and_a_label_over_a_figure_go_through_pdf(self):
        self.assertEqual(self.judge(arrows()), (True, "drawings over 1 of its 1 figures"))

    def test_a_label_from_the_next_paragraph(self):
        self.assertTrue(self.judge(label_after())[0])

    def test_a_group_and_a_canvas(self):
        self.assertEqual(self.judge(grouped("wpg")), (True, "drawings over 1 of its 1 figures"))
        self.assertEqual(self.judge(grouped("wpc")), (True, "drawings over 1 of its 1 figures"))

    def test_vml(self):
        self.assertEqual(self.judge(vml()), (True, "drawings over 1 of its 1 figures"))

    def test_a_frame_is_not_a_drawing_and_the_fallback_is_not_read(self):
        self.assertEqual(self.judge(alternate()), (False, "no drawings over its figure"))

    def test_a_wrapped_text_box_is_not_over_the_figure(self):
        self.assertFalse(self.judge(sidebar())[0])

    def test_a_manual_with_callouts_on_a_few_figures_reads_native(self):
        self.assertEqual(self.judge(manual()), (False, "drawings over 1 of its 6 figures, under a third"))

    def test_icons_are_not_figures(self):
        self.assertEqual(self.judge(icons()), (False, "no figures"))

    def test_a_broken_docx_goes_through_pdf(self):
        pdf, reason = self.judge(b"PK\x03\x04 not a zip")
        self.assertTrue(pdf)
        self.assertIn("could not be read", reason)


class RouteTest(unittest.TestCase):
    """The way call_tool sends a document, LibreOffice and docling stubbed."""

    def setUp(self):
        self.calls = []

        def office_convert(path, into, target):
            self.calls.append((path.suffix, target))
            out = pathlib.Path(into) / "office"
            out.mkdir(exist_ok=True)
            got = out / (path.stem + "." + target)
            if target == "docx":
                got.write_bytes(self.as_docx)
            else:
                got.write_bytes(b"%PDF-1.4 stub")
            return got

        def convert_sync(name, source, *args):
            self.read = pathlib.Path(source).suffix if not isinstance(source, bytes) else "bytes"
            self.args = args
            return {"markdown": "x\n", "title": "", "pages": 1, "tables": 0, "pictures": []}, []

        for target, value in (("SOFFICE", "/usr/bin/soffice"), ("office_convert", office_convert), ("convert_sync", convert_sync)):
            patcher = mock.patch.object(server, target, value)
            patcher.start()
            self.addCleanup(patcher.stop)
        # The kinds LibreOffice reads, as when it is there.
        patcher = mock.patch.dict(server.KINDS, {ext: mime for ext, (_, mime) in server.OFFICE.items()})
        patcher.start()
        self.addCleanup(patcher.stop)

    def call(self, name, data, **settings):
        self.as_docx = data
        r = asyncio.run(server.call_tool("convert", {"file": {"name": name, "data": base64.b64encode(data).decode()}} | settings))
        self.assertFalse(r.isError, r.content)
        return r.structuredContent

    def test_auto_reads_a_plain_docx_native(self):
        got = self.call("care.docx", plain())
        self.assertEqual((got["office_used"], got["office_reason"]), ("native", "no drawings over its 2 figures"))
        self.assertEqual(self.calls, [])
        self.assertEqual(self.read, ".docx")

    def test_auto_sends_a_drawn_docx_through_pdf(self):
        got = self.call("coulter.docx", arrows())
        self.assertEqual((got["office_used"], got["office_reason"]), ("pdf", "drawings over 1 of its 1 figures"))
        self.assertEqual(self.calls, [(".docx", "pdf")])
        self.assertEqual(self.read, ".pdf")

    def test_auto_turns_a_doc_into_docx_first(self):
        got = self.call("coulter.doc", arrows())
        self.assertEqual(got["office_used"], "pdf")
        self.assertEqual(self.calls, [(".doc", "docx"), (".doc", "pdf")])
        got = self.call("care.rtf", plain())
        self.assertEqual(got["office_used"], "native")
        self.assertEqual(self.calls[2:], [(".rtf", "docx")])
        self.assertEqual(self.read, ".docx")

    def test_forced(self):
        got = self.call("care.docx", plain(), office="pdf")
        self.assertEqual((got["office_used"], got["office_reason"]), ("pdf", "office: pdf"))
        got = self.call("coulter.docx", arrows(), office="native")
        self.assertEqual((got["office_used"], got["office_reason"]), ("native", "office: native"))
        self.assertEqual(self.read, ".docx")

    def test_no_libreoffice(self):
        with mock.patch.object(server, "SOFFICE", None):
            got = self.call("coulter.docx", arrows())
        self.assertEqual((got["office_used"], got["office_reason"]), ("native", "no LibreOffice to make a PDF"))

    def test_presentations_and_spreadsheets(self):
        got = self.call("deck.pptx", b"stub", slides=False)
        self.assertEqual((got["office_used"], got["office_reason"]), ("pdf", "a presentation without slides goes through PDF"))
        got = self.call("deck.odp", b"stub", slides=False, office="native")
        self.assertEqual(got["office_used"], "native")
        self.assertEqual(self.calls[-1], (".odp", "pptx"))
        got = self.call("deck.pptx", b"stub")
        self.assertNotIn("office_used", got)
        self.assertEqual(self.read, "bytes")
        got = self.call("sheet.ods", b"stub")
        self.assertNotIn("office_used", got)
        self.assertEqual(self.calls[-1], (".ods", "xlsx"))

    def test_picture_text_setting(self):
        self.call("care.docx", plain())
        self.assertTrue(self.args[-1])
        self.call("care.docx", plain(), picture_text=False)
        self.assertFalse(self.args[-1])

    def test_a_wrong_office(self):
        self.as_docx = plain()
        r = asyncio.run(server.call_tool("convert", {"file": {"name": "a.docx", "data": base64.b64encode(plain()).decode()}, "office": "word"}))
        self.assertTrue(r.isError)
        self.assertIn("auto, pdf, native", r.content[0].text)


class NativeTest(unittest.TestCase):
    """A plain docx through docling's Word reader, as office: auto sends it."""

    def test_read_native(self):
        from docling.datamodel.base_models import InputFormat
        from docling.document_converter import DocumentConverter

        with mock.patch.object(server, "SOFFICE", "/usr/bin/soffice"), \
                mock.patch.object(server, "converter", lambda *a, **k: DocumentConverter(allowed_formats=[InputFormat.DOCX])):
            r = asyncio.run(server.call_tool("convert", {"file": {"name": "care.docx", "data": base64.b64encode(plain()).decode()}}))
        self.assertFalse(r.isError, r.content)
        got = r.structuredContent
        self.assertEqual(got["office_used"], "native")
        self.assertRegex(got["markdown"], r"(?m)^#+ Seed drill care$")
        self.assertRegex(got["markdown"], r"\| *Hopper *\| *Empty it before winter\. *\|")
        self.assertEqual([p["name"] for p in got["pictures"]], ["fig-1.png", "fig-2.png"])


# The text inside a picture --------------------------------------------------

def page_doc(children, picture_box=(50, 100, 550, 700), caption=None):
    """A page of 600 x 800 as docling's PDF pipeline leaves it: a heading,
    a picture holding the text children (text, (left, top, right, bottom)
    from the top of the page), and a closing line."""
    from docling_core.types.doc import BoundingBox, CoordOrigin, DocItemLabel, DoclingDocument, ImageRef, ProvenanceItem, Size

    def prov(box, text=""):
        l, t, r, b = box
        return ProvenanceItem(page_no=1, charspan=(0, len(text)), bbox=BoundingBox(l=l, t=800 - t, r=r, b=800 - b, coord_origin=CoordOrigin.BOTTOMLEFT))

    doc = DoclingDocument(name="page")
    doc.add_page(page_no=1, size=Size(width=600, height=800))
    doc.add_heading("Warning signs", prov=prov((50, 40, 300, 60), "Warning signs"))
    pic = doc.add_picture(prov=prov(picture_box), image=ImageRef.from_pil(Image.new("RGB", (100, 120), "white"), dpi=72))
    if caption:
        item = doc.add_text(label=DocItemLabel.CAPTION, text=caption, parent=pic, prov=prov((50, 705, 300, 715), caption))
        pic.captions.append(item.get_ref())
    for text, box in children:
        doc.add_text(label=DocItemLabel.TEXT, text=text, parent=pic, prov=prov(box, text))
    doc.add_text(label=DocItemLabel.TEXT, text="Keep this page with the machine.", prov=prov((50, 720, 400, 735)))
    return doc


# Two columns of pictograms, each with its label on one to three lines
# (icons from 60 to 120 and from 300 to 360, labels beside them); a word of
# the left column cut into two pieces on its line.
SIGNS = [
    ("WEAR", (130, 110, 170, 120)), ("GLOVES", (175, 110, 230, 120)),
    ("READ THE", (130, 200, 200, 210)), ("MANUAL FIRST", (130, 212, 230, 222)),
    ("KEEP CHILDREN", (130, 300, 240, 310)), ("AWAY FROM THE", (130, 312, 245, 322)), ("MACHINE", (130, 324, 195, 334)),
    ("HOT", (370, 110, 400, 120)), ("SURFACE", (370, 122, 430, 132)),
    ("NO OPEN", (370, 200, 430, 210)), ("FLAMES", (370, 212, 425, 222)),
    ("1. TOXIC", (370, 300, 430, 310)),
]


class PictureTextTest(unittest.TestCase):
    def convert(self, doc, fmt="PDF", **kwargs):
        from docling.datamodel.base_models import InputFormat

        result = SimpleNamespace(document=doc, status=SimpleNamespace(name="SUCCESS"), errors=[], input=SimpleNamespace(format=InputFormat[fmt]))
        fake = SimpleNamespace(convert=lambda *a, **k: result)
        args = dict(pictures=True, picture_text=True) | kwargs
        with mock.patch.object(server, "converter", lambda *a, **k: fake):
            return server.convert_sync("page.pdf", b"%PDF", 1 << 20, False, args["pictures"], 10, 2.0, "pypdfium", None, args["picture_text"])

    def test_the_labels_under_the_picture(self):
        result, files = self.convert(page_doc(SIGNS))
        self.assertEqual(len(files), 1)
        self.assertEqual(result["markdown"], "\n\n".join([
            "## Warning signs",
            "![](fig-1.png)",
            "WEAR GLOVES",
            "READ THE MANUAL FIRST",
            "KEEP CHILDREN AWAY FROM THE MACHINE",
            "HOT SURFACE",
            "NO OPEN FLAMES",
            "1\\. TOXIC",
            "Keep this page with the machine.",
        ]) + "\n")

    def test_a_row_of_boxes_left_to_right(self):
        boxes = [("Fill", (60, 300, 100, 310)), ("the tank", (60, 312, 120, 322)), ("Start", (250, 300, 300, 310)),
                 ("the motor", (250, 312, 320, 322)), ("Spray", (450, 300, 500, 310)),
                 ("Stop the motor before cleaning the nozzles.", (60, 500, 500, 510))]
        md = self.convert(page_doc(boxes))[0]["markdown"]
        self.assertIn("![](fig-1.png)\n\nFill the tank\n\nStart the motor\n\nSpray\n\nStop the motor before cleaning the nozzles.\n\n", md)

    def test_right_to_left(self):
        # Two columns: a line in two pieces and a label on the right, two
        # labels on the left. The right one first, its line from the right.
        boxes = [("שלום", (460, 300, 500, 310)), ("עולם", (410, 300, 450, 310)), ("ימין", (440, 400, 480, 410)),
                 ("שמאל", (60, 300, 100, 310)), ("תחתון", (60, 400, 110, 410))]
        md = self.convert(page_doc(boxes))[0]["markdown"]
        self.assertIn("![](fig-1.png)\n\nשלום עולם\n\nימין\n\nשמאל\n\nתחתון\n\n", md)

    def test_a_split_word_and_the_caption(self):
        boxes = [("The pump draws water from the tank and sends it to the noz-", (60, 300, 500, 310)),
                 ("zles on the boom.", (60, 312, 200, 322))]
        md = self.convert(page_doc(boxes, caption="Figure 2: the pump"))[0]["markdown"]
        self.assertIn("Figure 2: the pump\n\n![Figure 2: the pump](fig-1.png)\n\nThe pump draws water from the tank and sends it to the nozzles on the boom.\n\n", md)
        self.assertEqual(md.count("Figure 2: the pump"), 2)  # the caption and the picture's alt

    def test_a_photo_with_a_stray_letter_keeps_none(self):
        md = self.convert(page_doc([("A", (200, 300, 210, 310)), ("STOP", (300, 300, 340, 310)), ("12", (300, 400, 320, 410))]))[0]["markdown"]
        self.assertIn("![](fig-1.png)\n\nKeep this page", md)

    def test_numbers_alone_are_left_out_and_a_chart_keeps_none(self):
        callouts = [("Drive belt", (60, 300, 130, 310)), ("1", (200, 300, 206, 310)), ("Belt guard", (60, 400, 130, 410)),
                    ("2", (200, 400, 206, 410)), ("Tension spring", (60, 500, 150, 510))]
        md = self.convert(page_doc(callouts))[0]["markdown"]
        self.assertIn("![](fig-1.png)\n\nDrive belt\n\nBelt guard\n\nTension spring\n\nKeep", md)
        chart = [(str(v), (100 + 40 * i, 600 - 5 * v, 120 + 40 * i, 610 - 5 * v)) for i, v in enumerate((12, 30, 45, 61, 70))]
        chart += [("Sales by month, in crates", (100, 120, 400, 130)), ("2025", (100, 650, 140, 660))]
        md = self.convert(page_doc(chart))[0]["markdown"]
        self.assertIn("![](fig-1.png)\n\nKeep", md)

    def test_off(self):
        md = self.convert(page_doc(SIGNS), picture_text=False)[0]["markdown"]
        self.assertNotIn("GLOVES", md)

    def test_without_the_picture_files(self):
        result, files = self.convert(page_doc(SIGNS), pictures=False)
        self.assertEqual(files, [])
        self.assertIn("## Warning signs\n\nWEAR GLOVES\n\n", result["markdown"])

    def test_not_for_a_word_document(self):
        md = self.convert(page_doc(SIGNS), fmt="DOCX")[0]["markdown"]
        self.assertNotIn("GLOVES", md)

    def test_markdown_markers_escaped(self):
        self.assertEqual(server.md_line("# 3 bar"), "\\# 3 bar")
        self.assertEqual(server.md_line("- 5 °C"), "\\- 5 °C")
        self.assertEqual(server.md_line("2) Lift"), "2\\) Lift")
        self.assertEqual(server.md_line("---"), "\\---")
        self.assertEqual(server.md_line("<tap> A_B"), "&lt;tap&gt; A\\_B")


if __name__ == "__main__":
    unittest.main()
